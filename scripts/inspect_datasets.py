from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

DATA_DIR = Path("data")


def clean(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.replace("\xa0", " ")
        value = value.strip()

        return value if value else None

    return value


def inspect_sheet(sheet):
    rows = list(sheet.iter_rows(values_only=True))

    if len(rows) < 2:
        return

    headers = [clean(value) for value in rows[1]]

    data_rows = rows[2:]

    print(f"\n{'=' * 70}")
    print(f"SHEET: {sheet.title}")
    print(f"{'=' * 70}")

    print(f"Data rows: {len(data_rows)}")
    print(f"Columns: {len(headers)}")

    print("\nCOLUMNS:")
    for index, header in enumerate(headers, start=1):
        print(f"  {index}. {header}")

    print("\nMISSING VALUES:")

    for column_index, header in enumerate(headers):
        values = [
            clean(row[column_index])
            for row in data_rows
            if column_index < len(row)
        ]

        missing = sum(value is None for value in values)

        print(
            f"  {header}: "
            f"{missing}/{len(data_rows)} missing"
        )

    # Look for duplicates in likely identity fields.
    print("\nDUPLICATES:")

    for field_name in [
        "NBFC Name",
        "Name of the company",
        "Corporate Identification Number",
        "NBFC Code",
    ]:
        if field_name not in headers:
            continue

        index = headers.index(field_name)

        values = [
            clean(row[index])
            for row in data_rows
            if index < len(row) and clean(row[index]) is not None
        ]

        counts = Counter(values)

        duplicates = [
            (value, count)
            for value, count in counts.items()
            if count > 1
        ]

        print(
            f"  {field_name}: "
            f"{len(duplicates)} duplicated values"
        )

        for value, count in duplicates[:10]:
            print(f"      {count}x {value}")

    # Useful categorical fields.
    print("\nUNIQUE VALUES:")

    for field_name in [
        "Classification",
        "Layer",
        "Whether have CoR for holding/ Accepting Public Deposits",
        "Regional Office",
        "Category",
        "Reason",
    ]:
        if field_name not in headers:
            continue

        index = headers.index(field_name)

        values = {
            clean(row[index])
            for row in data_rows
            if index < len(row) and clean(row[index]) is not None
        }

        print(f"\n  {field_name}: {len(values)} unique")

        for value in sorted(values, key=str)[:50]:
            print(f"      {value}")


def main():
    files = (
        list(DATA_DIR.glob("*.xlsx"))
        + list(DATA_DIR.glob("*.XLSX"))
    )

    if not files:
        print("No Excel files found in data/")
        return

    for file in files:

        print("\n\n")
        print("#" * 80)
        print(f"WORKBOOK: {file.name}")
        print("#" * 80)

        workbook = load_workbook(
            filename=file,
            read_only=True,
            data_only=True,
        )

        print(f"\nSheets: {workbook.sheetnames}")

        for sheet in workbook.worksheets:
            inspect_sheet(sheet)

        workbook.close()


if __name__ == "__main__":
    main()